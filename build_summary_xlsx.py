#!/usr/bin/env python3
"""納品書台帳（納品書集計.xlsx の「データ」シート）から、見やすい集計表を作る。

台帳は読むだけで一切書き換えない。出力は同じフォルダの「納品書_集計表.xlsx」。
何度実行しても最新の台帳から作り直すので、毎晩の記帳のあとに呼べば常に最新になる。

  python3 build_summary_xlsx.py            # OneDrive の既定の場所を使う
  python3 build_summary_xlsx.py <台帳> <出力>
"""
import os
import re
import sys
import tempfile
from collections import defaultdict
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ONEDRIVE = os.path.expanduser("~/Library/CloudStorage/OneDrive-個人用")
LEDGER = os.path.join(ONEDRIVE, "納品書集計", "納品書集計.xlsx")
OUTPUT = os.path.join(ONEDRIVE, "納品書集計", "納品書_集計表.xlsx")
NAME_TABLE = os.path.join(ONEDRIVE, "納品書写真", "_仕入先の正式名称.txt")

STORE_ORDER = ["本店", "KADODE店", "空港店", "静岡紺屋町店", "セントラル", "製麺事業部", "冷凍事業部"]
# 正式名称表に無いが台帳に出てくる別表記（英字表記・運営会社名など）
EXTRA_ALIASES = {"MonotaRO": "モノタロウ", "世亜企画": "韓国市場"}

NAVY = "1F2A44"
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
HEAD_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="E8ECF3")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="C8CED8")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
YEN = '#,##0"円";-#,##0"円";""'


def load_name_table(path):
    """「正式名称 ← 別表記, 別表記」の表を {表記: 正式名称} にする。"""
    aliases = {}
    if not os.path.exists(path):
        return aliases
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            official, _, rest = line.partition("←")
            official = official.strip()
            aliases[official] = official
            for a in re.split(r"[,、]", rest):
                if a.strip():
                    aliases[a.strip()] = official
    return aliases


def normalize(name, aliases):
    """台帳に残っている表記ゆれを、集計表の上でだけ正式名称に揃える。"""
    name = (name or "").strip()
    if name in aliases:
        return aliases[name]
    bare = re.sub(r"(株式会社|有限会社|（株）|\(株\))", "", name).strip()
    # 「新村商店（株式会社シンムラCOMPANY）」「株式会社世亜企画（韓国市場）」のような括弧書き
    m = re.match(r"^(.*?)[（(](.*?)[）)]$", bare)
    candidates = [bare] + ([m.group(1).strip(), m.group(2).strip()] if m else [])
    for c in candidates:
        if c in aliases:
            return aliases[c]
        for key, official in EXTRA_ALIASES.items():
            if key in c:
                return official
    return candidates[1] if m and candidates[1] else bare


def read_ledger(path, aliases):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["データ"]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        date, store, supplier, amount = r[0], r[1], r[2], r[3]
        if isinstance(date, datetime):
            date = date.strftime("%Y-%m-%d")
        date = str(date)[:10]
        try:
            amount = int(amount or 0)
        except (TypeError, ValueError):
            amount = 0
        rows.append({
            "date": date,
            "month": date[:7],
            "store": (store or "不明").strip(),
            "supplier": normalize(supplier, aliases),
            "amount": amount,
            "kind": (r[5] if len(r) > 5 else None) or "",
        })
    wb.close()
    return rows


def store_key(s):
    return (STORE_ORDER.index(s) if s in STORE_ORDER else len(STORE_ORDER), s)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")


def title(ws, text, note):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = note
    ws["A2"].font = Font(color="666666", size=9)


def write_matrix(ws, top, row_label, row_keys, col_keys, values):
    """行×列のクロス表を書き、右端と最下行に合計を付ける。次に書ける行番号を返す。"""
    headers = [row_label] + list(col_keys) + ["合計"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=top, column=i, value=h)
    style_header(ws, top, len(headers))
    r = top + 1
    col_totals = defaultdict(int)
    for rk in row_keys:
        ws.cell(row=r, column=1, value=rk).border = BORDER
        line_total = 0
        for j, ck in enumerate(col_keys, 2):
            v = values.get((rk, ck), 0)
            cell = ws.cell(row=r, column=j, value=v or None)
            cell.number_format, cell.border = YEN, BORDER
            line_total += v
            col_totals[ck] += v
        cell = ws.cell(row=r, column=len(headers), value=line_total)
        cell.number_format, cell.border, cell.font = YEN, BORDER, BOLD
        r += 1
    ws.cell(row=r, column=1, value="合計")
    for j, ck in enumerate(col_keys, 2):
        ws.cell(row=r, column=j, value=col_totals[ck]).number_format = YEN
    ws.cell(row=r, column=len(headers), value=sum(col_totals.values())).number_format = YEN
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=r, column=c)
        cell.fill, cell.font, cell.border = TOTAL_FILL, BOLD, BORDER
    return r + 2


def set_widths(ws, first, rest, n):
    ws.column_dimensions["A"].width = first
    for c in range(2, n + 1):
        ws.column_dimensions[get_column_letter(c)].width = rest


def build(rows, out_path, ledger_path):
    months = sorted({r["month"] for r in rows})
    stores = sorted({r["store"] for r in rows}, key=store_key)
    by = lambda keyf: _sum(rows, keyf)
    sup_total = by(lambda r: r["supplier"])
    suppliers = sorted(sup_total, key=lambda s: -sup_total[s])
    stamp = f"{datetime.now():%Y-%m-%d %H:%M} 作成 ／ 元データ：{os.path.basename(ledger_path)}（{len(rows)}行）"

    wb = openpyxl.Workbook()

    # 1. 明細（絞り込み用）
    ws = wb.active
    ws.title = "明細（絞り込み）"
    title(ws, "納品書の明細", stamp + " ／ 見出しの ▼ で店舗・取引先・月を絞り込めます")
    ws["A3"], ws["C3"] = "表示中の件数", "表示中の合計"
    ws["A3"].font = ws["C3"].font = BOLD
    head = 5
    for i, h in enumerate(["日付", "年月", "店舗", "取引先", "金額", "区分"], 1):
        ws.cell(row=head, column=i, value=h)
    style_header(ws, head, 6)
    detail = sorted(rows, key=lambda r: (r["date"], store_key(r["store"])), reverse=True)
    for i, r in enumerate(detail, head + 1):
        for j, v in enumerate([r["date"], r["month"], r["store"], r["supplier"], r["amount"], r["kind"]], 1):
            ws.cell(row=i, column=j, value=v).border = BORDER
        ws.cell(row=i, column=5).number_format = YEN
    last = head + max(len(detail), 1)
    ws["B3"] = f"=SUBTOTAL(103,A{head + 1}:A{last})"
    ws["B3"].number_format = '#,##0"件"'
    ws["D3"] = f"=SUBTOTAL(109,E{head + 1}:E{last})"
    ws["D3"].number_format = YEN
    ws["D3"].font = Font(bold=True, size=12, color=NAVY)
    ws.auto_filter.ref = f"A{head}:F{last}"
    ws.freeze_panes = f"A{head + 1}"
    for col, w in zip("ABCDEF", [12, 10, 14, 20, 14, 10]):
        ws.column_dimensions[col].width = w

    # 2. 店舗×月
    ws = wb.create_sheet("店舗×月")
    title(ws, "店舗ごとの月別仕入れ", stamp)
    write_matrix(ws, 4, "店舗", stores, months, by(lambda r: (r["store"], r["month"])))
    set_widths(ws, 16, 14, len(months) + 2)
    ws.freeze_panes = "B5"

    # 3. 取引先×月
    ws = wb.create_sheet("取引先×月")
    title(ws, "取引先ごとの月別仕入れ（金額の多い順）", stamp)
    write_matrix(ws, 4, "取引先", suppliers, months, by(lambda r: (r["supplier"], r["month"])))
    set_widths(ws, 20, 14, len(months) + 2)
    ws.freeze_panes = "B5"

    # 4. 取引先×店舗
    ws = wb.create_sheet("取引先×店舗")
    title(ws, "取引先ごとの店舗別仕入れ（全期間・金額の多い順）", stamp)
    write_matrix(ws, 4, "取引先", suppliers, stores, by(lambda r: (r["supplier"], r["store"])))
    set_widths(ws, 20, 14, len(stores) + 2)
    ws.freeze_panes = "B5"

    # 5. 店舗ごとの内訳（店舗ごとに 取引先×月）
    ws = wb.create_sheet("店舗ごとの内訳")
    title(ws, "店舗ごとの取引先別・月別仕入れ", stamp)
    top = 4
    for s in stores:
        srows = [r for r in rows if r["store"] == s]
        tot = _sum(srows, lambda r: r["supplier"])
        ws.cell(row=top, column=1, value=f"■ {s}").font = Font(bold=True, size=12, color=NAVY)
        top = write_matrix(ws, top + 1, "取引先", sorted(tot, key=lambda k: -tot[k]), months,
                           _sum(srows, lambda r: (r["supplier"], r["month"])))
    set_widths(ws, 20, 14, len(months) + 2)

    # 6. 取引先ごとの内訳（取引先ごとに 店舗×月）
    ws = wb.create_sheet("取引先ごとの内訳")
    title(ws, "取引先ごとの店舗別・月別仕入れ（金額の多い順）", stamp)
    top = 4
    for sup in suppliers:
        srows = [r for r in rows if r["supplier"] == sup]
        ws.cell(row=top, column=1, value=f"■ {sup}").font = Font(bold=True, size=12, color=NAVY)
        top = write_matrix(ws, top + 1, "店舗", sorted({r["store"] for r in srows}, key=store_key), months,
                           _sum(srows, lambda r: (r["store"], r["month"])))
    set_widths(ws, 20, 14, len(months) + 2)

    # 途中で止まっても壊れたファイルを残さないよう、一時ファイルに書いてから置き換える
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(out_path))
    os.close(fd)
    wb.save(tmp)
    os.replace(tmp, out_path)


def _sum(rows, keyf):
    d = defaultdict(int)
    for r in rows:
        d[keyf(r)] += r["amount"]
    return d


def main():
    ledger = sys.argv[1] if len(sys.argv) > 1 else LEDGER
    out = sys.argv[2] if len(sys.argv) > 2 else OUTPUT
    aliases = load_name_table(NAME_TABLE)
    rows = read_ledger(ledger, aliases)
    build(rows, out, ledger)
    unknown = sorted({r["supplier"] for r in rows} - set(aliases.values()) - set(EXTRA_ALIASES.values()))
    print(f"作成しました: {out}（{len(rows)}行 / 合計 {sum(r['amount'] for r in rows):,}円）")
    if unknown:
        print("正式名称表に無い取引先:", "、".join(unknown))


if __name__ == "__main__":
    main()
