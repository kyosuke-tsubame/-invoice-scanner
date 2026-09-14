#!/usr/bin/env python3
"""台帳（納品書集計.xlsx）の「仕入先」列を、名寄せ表の正式名称に揃える（1回きりの手直し用）。

2026-09-15作成。確認画面の修正前などに旧表記（株式会社クリチク 等）で記帳された行を直す。
日付・店舗・金額・年月・区分は一切変えない。実行前に .bak.<日時>_名寄せ前 のバックアップを作り、
書き換え後に「名前以外がバックアップと全く同じか」を検証して表示する。

  python3 normalize_ledger_names.py          # 変わる行を表示するだけ（書き換えない）
  python3 normalize_ledger_names.py --apply  # バックアップを取ってから書き換える
"""
import shutil
import sys
import time
from collections import Counter

from openpyxl import load_workbook

import invoice_ledger as L


def data_rows(path):
    with open(path, "rb") as f:
        ws = load_workbook(f, read_only=True)[L.DATA_SHEET]
        return [r for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0] is not None]


def main():
    apply = "--apply" in sys.argv
    wb = load_workbook(L.XLSX_PATH)
    ws = wb[L.DATA_SHEET]
    changes = Counter()
    for row in ws.iter_rows(min_row=2):
        if row[0].value is None:
            continue
        new = L.normalize_supplier(row[2].value)
        if new != row[2].value:
            changes[(row[2].value, new)] += 1
            row[2].value = new
    for (old, new), n in sorted(changes.items()):
        print(f"  {old} → {new} ×{n}")
    print(f"書き換わる行: {sum(changes.values())}")
    if not apply or not changes:
        print("（書き換えていません。実行するには --apply を付けてください）" if changes else "揃っていない行はありません。")
        return

    backup = f"{L.XLSX_PATH}.bak.{time.strftime('%Y%m%d_%H%M%S')}_名寄せ前"
    shutil.copy2(L.XLSX_PATH, backup)
    L._rebuild_summary(wb, L._read_rows(ws))
    wb.save(L.XLSX_PATH)

    a, b = data_rows(backup), data_rows(L.XLSX_PATH)
    same = len(a) == len(b) and all(
        x[:2] + x[3:6] == y[:2] + y[3:6] and L.normalize_supplier(x[2]) == y[2] for x, y in zip(a, b)
    )
    print(f"バックアップ: {backup}")
    print(f"行数 {len(a)} → {len(b)} ／ 名前以外は全く同じ: {same}")
    print(f"合計金額 {sum(int(r[3] or 0) for r in a):,} → {sum(int(r[3] or 0) for r in b):,}")
    print(f"取引先 {len({r[2] for r in a})}通り → {len({r[2] for r in b})}通り")


if __name__ == "__main__":
    main()
