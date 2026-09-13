#!/usr/bin/env python3
"""要確認になった納品書を、1枚のExcelで直せるようにする。

2026-09-13新設。背景：これまで要確認は「Slackに件数が出る → 人がOneDriveで写真を探して
Excel台帳に手入力」という流れで、①作業が重い ②直しても内部の記録は要確認のままで
永久に滞留カウントに残る、という問題があった（実際に、直したのに残り続けた記録が16件あった）。

使い方:
  python3 review_sheet.py build   # 要確認の一覧をExcelに書き出す
  python3 review_sheet.py apply   # 記入されたExcelを読んで、台帳に記帳し要確認を解除する

置き場所: OneDrive/納品書集計/要確認シート.xlsx
"""
import json
import shutil
import sys
import time
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import invoice_ledger as L  # noqa: E402
import validate_entry as V  # noqa: E402

BASE = Path(__file__).resolve().parent
STATE_FILE = BASE / "invoice_ocr_state.json"
SHEET_PATH = Path.home() / "Library/CloudStorage/OneDrive-個人用/納品書集計/要確認シート.xlsx"
SHEET_NAME = "要確認"
PHOTOS_ROOT = Path.home() / "Library/CloudStorage/OneDrive-個人用/納品書写真"

HEADERS = [
    "状態", "店舗(読み取り)", "仕入先(読み取り)", "日付(読み取り)", "金額(読み取り)", "なぜ要確認か",
    "正しい店舗", "正しい仕入先", "正しい日付", "正しい金額(税抜)", "記帳する？", "今後も同じ扱い？",
    "写真の場所",
]
# 記入してほしい列（G〜L）。ここだけ色を変えて、どこを触ればよいか一目で分かるようにする
INPUT_COLS = range(7, 13)
WIDTHS = [8, 14, 18, 13, 14, 46, 14, 18, 13, 16, 12, 14, 60]

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", start_color="1C2340")
INPUT_FILL = PatternFill("solid", start_color="FFF6D9")
INPUT_HEAD_FILL = PatternFill("solid", start_color="B8860B")
DONE_FILL = PatternFill("solid", start_color="E8F5E9")
# 記入欄が「どこからどこまでが1マスか」分かるように、全セルに細い枠線を引く
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load_state():
    with open(STATE_FILE, encoding="utf-8") as f:
        return json.load(f)


def held_items(state):
    out = []
    for k, v in state.items():
        if not isinstance(v, dict) or v.get("status") != "held_for_review":
            continue
        if not Path(k).is_file():
            continue
        out.append((k, v))
    out.sort(key=lambda x: (x[1].get("store") or "", x[1].get("date") or "", x[1].get("supplier") or ""))
    return out


def cmd_build():
    state = load_state()
    items = held_items(state)

    # 既に記入済みの内容があれば引き継ぐ（作り直しで入力が消えないように）
    previous = {}
    if SHEET_PATH.is_file():
        try:
            old = load_workbook(SHEET_PATH)[SHEET_NAME]
            for row in old.iter_rows(min_row=2, values_only=True):
                if row and row[12]:
                    previous[str(row[12])] = row[6:12]
        except Exception:
            previous = {}

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    ws["A1"] = (
        "要確認の納品書一覧です。黄色い欄（G〜L列）だけ記入してください。"
        "読み取りが合っていれば「正しい〜」欄は空のままでOK、「記帳する？」に『はい』だけ入れてください。"
    )
    ws["A1"].font = Font(bold=True, size=12)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws.row_dimensions[1].height = 30

    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.font = HEAD_FONT
        cell.fill = INPUT_HEAD_FILL if c in INPUT_COLS else HEAD_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = BORDER

    for i, (path, v) in enumerate(items, start=3):
        prev = previous.get(path, (None,) * 6)
        ws.cell(row=i, column=1, value="未対応")
        ws.cell(row=i, column=2, value=v.get("store"))
        ws.cell(row=i, column=3, value=v.get("supplier"))
        ws.cell(row=i, column=4, value=v.get("date"))
        cell = ws.cell(row=i, column=5, value=v.get("total"))
        cell.number_format = "#,##0"
        note = (v.get("note") or "").strip()
        ws.cell(row=i, column=6, value=note).alignment = Alignment(wrap_text=True, vertical="top")
        for j, col in enumerate(INPUT_COLS):
            c = ws.cell(row=i, column=col, value=prev[j])
            c.fill = INPUT_FILL
            c.alignment = Alignment(horizontal="center", vertical="center")
            if col == 10:
                c.number_format = "#,##0"
                c.alignment = Alignment(horizontal="right", vertical="center")
        link = ws.cell(row=i, column=13, value=path)
        link.hyperlink = Path(path).as_uri()
        link.font = Font(color="0563C1", underline="single", size=9)
        for col in range(1, len(HEADERS) + 1):
            ws.cell(row=i, column=col).border = BORDER

    last = max(len(items) + 2, 3)
    dv_store = DataValidation(type="list", formula1='"' + ",".join(V.STORES) + '"', allow_blank=True)
    dv_yn = DataValidation(type="list", formula1='"はい,いいえ"', allow_blank=True)
    ws.add_data_validation(dv_store)
    ws.add_data_validation(dv_yn)
    dv_store.add(f"G3:G{last}")
    dv_yn.add(f"K3:K{last}")
    dv_yn.add(f"L3:L{last}")

    for c, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:M{last}"

    SHEET_PATH.parent.mkdir(parents=True, exist_ok=True)
    if SHEET_PATH.is_file():
        shutil.copy2(SHEET_PATH, SHEET_PATH.with_name(
            SHEET_PATH.name + ".bak." + time.strftime("%Y%m%d_%H%M%S")))
    wb.save(SHEET_PATH)
    print(f"作成しました: {SHEET_PATH}")
    print(f"  要確認 {len(items)}件を書き出しました（記入済みの内容は引き継ぎました）")
    return 0


def _norm_date(x):
    if x is None or str(x).strip() == "":
        return None
    if isinstance(x, (datetime, date)):
        return x.strftime("%Y-%m-%d")
    return str(x).strip()[:10].replace("/", "-")


def cmd_apply():
    if not SHEET_PATH.is_file():
        print(f"確認シートがありません: {SHEET_PATH}")
        print("先に python3 review_sheet.py build を実行してください。")
        return 1

    ws = load_workbook(SHEET_PATH)[SHEET_NAME]
    state = load_state()
    shutil.copy2(STATE_FILE, str(STATE_FILE) + ".bak." + time.strftime("%Y%m%d_%H%M%S") + "_確認シート適用前")

    added, skipped, untouched, errors, learn = [], [], 0, [], []
    rows_cache = V.ledger_rows()

    for row in ws.iter_rows(min_row=3):
        path = row[12].value
        if not path:
            continue
        want = (str(row[10].value or "")).strip()
        if want == "":
            untouched += 1
            continue
        v = state.get(path)
        if not isinstance(v, dict):
            errors.append(f"{Path(path).name}: 内部の記録が見つかりません")
            continue

        entry = {
            "date": _norm_date(row[8].value) or v.get("date"),
            "store": (str(row[6].value).strip() if row[6].value else None) or v.get("store"),
            "supplier": (str(row[7].value).strip() if row[7].value else None) or v.get("supplier"),
            "total": int(row[9].value) if row[9].value not in (None, "") else v.get("total"),
        }

        if want == "いいえ":
            v["status"] = "skipped"
            v["resolvedAt"] = time.strftime("%Y-%m-%dT%H:%M:%S+09:00")
            v["resolvedNote"] = "確認シートで『記帳しない』と指定された"
            skipped.append(f"{entry['store']} / {entry['supplier']} / {entry['date']}")
            continue
        if want != "はい":
            errors.append(f"{Path(path).name}: 『記帳する？』は「はい」か「いいえ」で入れてください（今は『{want}』）")
            continue

        # 「はい」でも、記帳に必要な項目が埋まっていなければ書かない。
        # 日付が空のまま台帳に入ると、行はあるのに集計から漏れる幽霊行になるため。
        missing = []
        if not entry["date"] or V.parse_date(entry["date"]) is None:
            missing.append("正しい日付")
        if entry["store"] not in V.STORES:
            missing.append("正しい店舗")
        if not entry["supplier"]:
            missing.append("正しい仕入先")
        if not isinstance(entry["total"], int) or entry["total"] <= 0:
            missing.append("正しい金額(税抜)")
        if missing:
            errors.append(
                f"{Path(path).name}: 読み取れていない項目があるので記帳できません。"
                f"{'・'.join(missing)} を記入してください"
            )
            continue

        level, reasons = V.check(entry, rows=rows_cache)
        L.cmd_add({**entry, "category": "仕入れ"})
        rows_cache = V.ledger_rows()
        v.update({
            "status": "auto_saved",
            "store": entry["store"], "supplier": entry["supplier"],
            "date": entry["date"], "total": entry["total"],
            "resolvedAt": time.strftime("%Y-%m-%dT%H:%M:%S+09:00"),
            "resolvedNote": "確認シートで人が確認して記帳した"
            + (f"（検算の指摘: {'／'.join(reasons)}）" if reasons else ""),
        })
        added.append(f"{entry['store']} / {entry['supplier']} / {entry['date']} / {entry['total']:,}円"
                     + (f"  ※検算の指摘あり: {reasons[0]}" if reasons else ""))
        if str(row[11].value or "").strip() == "はい":
            learn.append((entry["supplier"], entry["store"]))

    with open(STATE_FILE, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False, indent=2)

    print(f"記帳した: {len(added)}件 ／ 記帳しないと指定: {len(skipped)}件 ／ 未記入のまま: {untouched}件")
    for x in added:
        print("  記帳 ", x)
    for x in skipped:
        print("  対象外", x)
    if errors:
        print("\n直してほしい記入ミス:")
        for e in errors:
            print("  -", e)
    if learn:
        print("\n『今後も同じ扱い』が指定されたもの（ルールへの反映は次の段階で実装）:")
        for sup, store in learn:
            print(f"  - {sup} → {store}")
    return 0


def main():
    if len(sys.argv) != 2 or sys.argv[1] not in ("build", "apply"):
        print(__doc__)
        return 1
    return cmd_build() if sys.argv[1] == "build" else cmd_apply()


if __name__ == "__main__":
    sys.exit(main())
